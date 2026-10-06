"""
Parsers de los reportes anuales de la SEPS (cooperativas S1-S3 + mutualistas).

Tres archivos por año (ver docs/fuentes_datos.md sección 4.0 y src/benchmark_bancos/config/
SEPS_DOWNLOAD_IDS), cada uno conformando contra tablas ya existentes:

- captaciones (ZIP de .xlsm/.xlsx, hoja `Base_captaciones`) -> staging.depositos
- colocaciones (ZIP de .xlsm, hojas `Base_colocaciones*`) -> staging.cartera. A pesar
  del nombre son SALDOS por estado de morosidad (por vencer / no devenga / vencida),
  no volumen de desembolsos.
- eeff (ZIP con un .txt delimitado) -> staging.boletin_balance / staging.boletin_pyg

Variaciones reales entre años que este módulo absorbe (verificadas 2026-09-30 contra
2021-2025, no supuestas):
- ZIPs con Deflate64 (zipfile de la stdlib no los abre) -> stream_unzip.
- Colocaciones 2021: DOS juegos de archivos en el mismo ZIP (`abr_2021` = ene-abr con la
  segmentación de crédito previa a la reforma JPRF de mayo 2021: CONSUMO PRIORITARIO/
  ORDINARIO, COMERCIAL PRIORITARIO/ORDINARIO; `dic_2021` = may-dic).
- Colocaciones S1 2025: partido por semestre en 2 hojas (`Base_colocaciones` = jul-dic,
  `Base_colocacionesISEM` = ene-jun, con FECHA DE CORTE como serial de Excel).
- Hoja `Base_para_actual`: auxiliar de la portada (fecha fija 2017-06-30), se ignora.
- Filas vacías al final de algunas hojas (RUC None).
- EEFF 2021: separador ';', encabezados con '_' y BOM; 2022+: TAB con comillas. Decimal
  '.' hasta ~2024 y ',' en 2025.
"""

import calendar
import datetime
import io
import logging
import re
import shutil
import unicodedata
from pathlib import Path

import openpyxl
import pandas as pd
from stream_unzip import stream_unzip

from benchmark_bancos.config import SEPS_RUC_SEGUNDO_PISO
from benchmark_bancos.transform.banco_matching import resolver_entidad_seps
from benchmark_bancos.transform.canton_matching import resolver_canton_bce
from benchmark_bancos.transform.categoria_deposito_matching import (
    resolver_categoria_deposito_seps,
)
from benchmark_bancos.transform.common import region_for_provincia, sha256_file

log = logging.getLogger(__name__)

_EXCEL_EPOCH = datetime.date(1899, 12, 30)

# SUBTIPO DE CREDITO (SEPS) -> tipo_credito canónico de CAPCOL (el mismo vocabulario que
# _REFRESH_MARTS_SQL ya mapea a marts.dim_segmento_credito). Incluye la segmentación
# previa a mayo 2021 (PRIORITARIO/ORDINARIO) colapsada a su segmento grueso actual, igual
# que dim_segmento_credito agrupa los sub-segmentos BCE. OPERACIONES CONTINGENTES no es
# cartera de la cuenta 14 -> se excluye (None).
SUBTIPO_CREDITO_SEPS = {
    "CONSUMO": "consumo",
    "CONSUMO PRIORITARIO": "consumo",
    "CONSUMO ORDINARIO": "consumo",
    "MICROCREDITO": "microcredito",
    "PRODUCTIVO": "comercial",
    "COMERCIAL PRIORITARIO": "comercial",
    "COMERCIAL ORDINARIO": "comercial",
    "INMOBILARIO": "inmobiliario",  # sic en la fuente
    "INMOBILIARIO": "inmobiliario",
    "VIVIENDA DE INTERES SOCIAL Y PUBLICO": "vivienda_interes_publico",
    "VIVIENDA INTERES PUBLICO Y SOCIAL": "vivienda_interes_publico",
    "EDUCATIVO": "educativo",
    "OPERACIONES CONTINGENTES": None,
}

ESTADOS_CARTERA_SEPS = {
    "CARTERA POR VENCER": "por_vencer",
    "CARTERA QUE NO DEVENGA INTERESES": "no_devenga_intereses",
    "CARTERA VENCIDA": "vencida",
}


class SubtipoCreditoSepsNoMapeadoError(ValueError):
    """Valor nuevo de SUBTIPO DE CREDITO -- agregarlo a SUBTIPO_CREDITO_SEPS."""


def _sin_tildes(texto: str) -> str:
    return "".join(
        c for c in unicodedata.normalize("NFKD", texto) if not unicodedata.combining(c)
    )


def _norm_header(h) -> str | None:
    if h is None:
        return None
    return " ".join(_sin_tildes(str(h)).replace("_", " ").upper().split())


def _chunks(path: Path, size: int = 1 << 20):
    with open(path, "rb") as f:
        while b := f.read(size):
            yield b


def extraer_zip_seps(
    zip_path: Path, extract_dir: Path, extensiones: tuple[str, ...]
) -> list[Path]:
    """Extrae (streaming, soporta Deflate64) los miembros con esas extensiones a
    extract_dir/<stem del zip>/ y devuelve sus rutas, ordenadas."""
    dest_dir = extract_dir / zip_path.stem
    if dest_dir.exists():
        shutil.rmtree(dest_dir)
    dest_dir.mkdir(parents=True)
    salida = []
    for name, _size, chunks in stream_unzip(_chunks(zip_path)):
        nombre = name.decode("utf-8", "replace")
        if not nombre.lower().endswith(extensiones):
            for _ in chunks:  # stream_unzip exige consumir cada miembro
                pass
            continue
        dest = dest_dir / Path(nombre).name
        with open(dest, "wb") as f:
            for c in chunks:
                f.write(c)
        salida.append(dest)
    if not salida:
        raise ValueError(
            f"{zip_path.name}: ningún archivo {extensiones} dentro del zip"
        )
    return sorted(salida)


def tipo_entidad_desde_archivo(nombre_archivo: str) -> str:
    """Boletin_captaciones_Dic25_Mut.xlsm / Reporte_colocaciones_dic_2025_MUT.xlsm ->
    MUTUALISTA; S1/S2/S3/SG1 -> COOPERATIVA. La SEPS separa por archivo, no por columna.
    """
    stem = Path(nombre_archivo).stem.upper()
    if re.search(r"(^|_)MUT(_|$)", stem):
        return "MUTUALISTA"
    if re.search(r"(^|_)S[GE]?[123](_|$)", stem):
        return "COOPERATIVA"
    raise ValueError(
        f"No se pudo derivar tipo_entidad del archivo SEPS '{nombre_archivo}'"
    )


def _tipo_entidad(ruc: str, por_archivo: str) -> str:
    return "ENTIDAD DE SEGUNDO PISO" if ruc in SEPS_RUC_SEGUNDO_PISO else por_archivo


def month_end_date(d: datetime.date) -> datetime.date:
    """Fin de mes real (las fechas de corte SEPS ya lo son; esto lo garantiza)."""
    return datetime.date(d.year, d.month, calendar.monthrange(d.year, d.month)[1])


def _fecha_corte(valor) -> datetime.date:
    if isinstance(valor, datetime.datetime):
        return month_end_date(valor.date())
    if isinstance(valor, datetime.date):
        return month_end_date(valor)
    if isinstance(valor, (int, float)):
        return month_end_date(_EXCEL_EPOCH + datetime.timedelta(days=int(valor)))
    return month_end_date(pd.to_datetime(str(valor)).date())


def _ruc_str(ruc) -> str:
    return (str(int(ruc)) if isinstance(ruc, (int, float)) else str(ruc).strip()).zfill(
        13
    )


def _iter_base_rows(xlsx_path: Path, prefijo_hoja: str):
    """Recorre todas las hojas cuyo nombre empieza con `prefijo_hoja` (p.ej.
    Base_colocaciones + Base_colocacionesISEM), devolviendo dicts con encabezado
    normalizado. Salta filas sin RUC (relleno vacío al final de la hoja)."""
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        hojas = [s for s in wb.sheetnames if s.lower().startswith(prefijo_hoja.lower())]
        if not hojas:
            raise ValueError(
                f"{xlsx_path.name}: no hay hoja '{prefijo_hoja}*' ({wb.sheetnames})"
            )
        for hoja in hojas:
            rows = wb[hoja].iter_rows(values_only=True)
            header = [_norm_header(h) for h in next(rows)]
            idx = {h: i for i, h in enumerate(header) if h}
            if "RUC" not in idx:
                raise ValueError(f"{xlsx_path.name}/{hoja}: sin columna RUC ({header})")
            for r in rows:
                if r[idx["RUC"]] is None:
                    continue
                yield {h: (r[i] if i < len(r) else None) for h, i in idx.items()}
    finally:
        wb.close()


def _geo(provincia_cruda, canton_crudo) -> tuple[str, str, str]:
    canton, provincia = resolver_canton_bce(canton_crudo, provincia_cruda)
    return region_for_provincia(provincia), provincia, canton


def parse_seps_captaciones_file(
    zip_path: Path, extract_dir: Path
) -> tuple[pd.DataFrame, list]:
    """-> (df con las columnas de staging.depositos, entidades para
    upsert_banco_maestro_ruc). Suma sobre ESTADO OPERACIÓN (NUEVA/VIGENTE/RENOVADA
    particionan el stock: concilia con EEFF cuenta 21 a 0% de mediana)."""
    source_hash = sha256_file(zip_path)
    records, entidades = [], {}
    for xlsx in extraer_zip_seps(zip_path, extract_dir, (".xlsm", ".xlsx")):
        tipo_archivo = tipo_entidad_desde_archivo(xlsx.name)
        for row in _iter_base_rows(xlsx, "Base_captaciones"):
            ruc = _ruc_str(row["RUC"])
            codigo, nombre, tipo, ruc = resolver_entidad_seps(
                row["RAZON SOCIAL"], ruc, _tipo_entidad(ruc, tipo_archivo)
            )
            entidades[codigo] = (codigo, nombre, tipo, ruc)
            region, provincia, canton = _geo(row["PROVINCIA"], row["CANTON"])
            tipo_deposito = " ".join(str(row["TIPO DE DEPOSITO"]).upper().split())
            records.append(
                {
                    "fecha": _fecha_corte(row["FECHA DE CORTE"]),
                    "tipo_entidad": tipo,
                    "banco": f"{nombre} ({ruc})",
                    "banco_codigo": codigo,
                    "region": region,
                    "provincia": provincia,
                    "canton": canton,
                    "tipo_deposito": tipo_deposito,
                    "categoria_deposito": resolver_categoria_deposito_seps(
                        tipo_deposito
                    ),
                    "saldo": float(row["SALDO"] or 0),
                    "numero_clientes": int(row["NUMERO DE CLIENTES"] or 0),
                    "numero_cuentas": int(row["NUMERO DE CUENTAS"] or 0),
                }
            )
    shutil.rmtree(extract_dir / zip_path.stem, ignore_errors=True)

    df = pd.DataFrame.from_records(records)
    key = [
        "fecha",
        "tipo_entidad",
        "banco",
        "banco_codigo",
        "region",
        "provincia",
        "canton",
        "tipo_deposito",
    ]
    df = df.groupby(key, as_index=False).agg(
        saldo=("saldo", "sum"),
        numero_clientes=("numero_clientes", "sum"),
        numero_cuentas=("numero_cuentas", "sum"),
        categoria_deposito=("categoria_deposito", "first"),
    )
    df["plazo_dias_desde"] = None  # el reporte SEPS con entidad no trae banda de plazo
    df["plazo_dias_hasta"] = None
    df["source_file"] = zip_path.name
    df["source_hash"] = source_hash
    return df, list(entidades.values())


def parse_seps_colocaciones_file(
    zip_path: Path, extract_dir: Path
) -> tuple[pd.DataFrame, list]:
    """-> (df en formato largo de staging.cartera, entidades). Agrega fuera los atributos
    que fact_saldo_cartera no modela (origen/estado/clase de operación, actividad
    económica) -- suman al grano (fecha, entidad, cantón, tipo_credito, estado_cartera).
    """
    source_hash = sha256_file(zip_path)
    agg: dict[tuple, float] = {}
    entidades = {}
    excluidas = 0.0
    for xlsx in extraer_zip_seps(zip_path, extract_dir, (".xlsm", ".xlsx")):
        tipo_archivo = tipo_entidad_desde_archivo(xlsx.name)
        n = 0
        for row in _iter_base_rows(xlsx, "Base_colocaciones"):
            n += 1
            subtipo = _sin_tildes(
                " ".join(str(row["SUBTIPO DE CREDITO"]).upper().split())
            )
            if subtipo not in SUBTIPO_CREDITO_SEPS:
                raise SubtipoCreditoSepsNoMapeadoError(
                    f"{xlsx.name}: SUBTIPO DE CREDITO no mapeado '{row['SUBTIPO DE CREDITO']}'. "
                    f"Agregarlo a SUBTIPO_CREDITO_SEPS en src/benchmark_bancos/transform/parse_seps.py."
                )
            tipo_credito = SUBTIPO_CREDITO_SEPS[subtipo]
            if tipo_credito is None:
                excluidas += float(row["CARTERA TOTAL"] or 0)
                continue
            ruc = _ruc_str(row["RUC"])
            codigo, nombre, tipo, ruc = resolver_entidad_seps(
                row["RAZON SOCIAL"], ruc, _tipo_entidad(ruc, tipo_archivo)
            )
            entidades[codigo] = (codigo, nombre, tipo, ruc)
            region, provincia, canton = _geo(row["PROVINCIA"], row["CANTON"])
            fecha = _fecha_corte(row["FECHA DE CORTE"])
            for col, estado in ESTADOS_CARTERA_SEPS.items():
                k = (
                    fecha,
                    tipo,
                    f"{nombre} ({ruc})",
                    codigo,
                    region,
                    provincia,
                    canton,
                    tipo_credito,
                    estado,
                )
                agg[k] = agg.get(k, 0.0) + float(row[col] or 0)
        log.info("SEPS colocaciones %s: %d filas leídas", xlsx.name, n)
    shutil.rmtree(extract_dir / zip_path.stem, ignore_errors=True)
    if excluidas:
        log.info(
            "SEPS colocaciones %s: %.2f USD de OPERACIONES CONTINGENTES excluidos",
            zip_path.name,
            excluidas,
        )

    cols = [
        "fecha",
        "tipo_entidad",
        "banco",
        "banco_codigo",
        "region",
        "provincia",
        "canton",
        "tipo_credito",
        "estado_cartera",
    ]
    df = pd.DataFrame([(*k, v) for k, v in agg.items()], columns=[*cols, "saldo"])
    df["source_file"] = zip_path.name
    df["source_hash"] = source_hash
    return df, list(entidades.values())


def _a_numero(serie: pd.Series) -> pd.Series:
    """Decimal '.' (2021-2024) o ',' (2025), sin separador de miles. Vacío -> NaN. Un
    valor no vacío que no parsea es un cambio de formato real -> fallo duro."""
    s = serie.str.strip().str.strip('"')
    vacio = s.eq("")
    valores = pd.to_numeric(s.str.replace(",", ".", regex=False), errors="coerce")
    malos = valores.isna() & ~vacio
    if malos.any():
        raise ValueError(
            f"EEFF SEPS: {malos.sum()} saldos no numéricos, ej. {s[malos].head(3).tolist()}"
        )
    return valores


def parse_seps_eeff_file(zip_path: Path, chunksize: int = 500_000) -> dict:
    """-> {'balance', 'pyg', 'cuentas', 'entidades'} con la forma que ya consume el flujo
    del Boletín (upsert_staging_boletin_balance/pyg). Cuentas 4* y 5* van a PYG, el resto
    a BALANCE. Se descartan saldos 0 o vacíos (~70% de las filas): la ausencia equivale a
    0 y así el volumen queda en ~900k filas/año en vez de ~3M. Valores ya en USD
    completos (a diferencia del Boletín, que viene en miles)."""
    source_hash = sha256_file(zip_path)
    gen = stream_unzip(_chunks(zip_path))
    _name, _size, member = next(gen)
    stream = io.BufferedReader(_IterStream(member), buffer_size=1 << 20)
    texto = io.TextIOWrapper(stream, encoding="utf-8-sig", newline="")
    primera = texto.readline()
    sep = ";" if primera.count(";") > primera.count("\t") else "\t"
    header = [
        _norm_header(h.strip().strip('"')) for h in primera.rstrip("\r\n").split(sep)
    ]
    renombre = {
        "FECHA DE CORTE": "fecha",
        "SEGMENTO": "segmento",
        "RUC": "ruc",
        "RAZON SOCIAL": "razon_social",
        "CUENTA": "codigo",
        "DESCRIPCION CUENTA": "cuenta",
        "SALDO (USD)": "saldo",
        "SALDO USD": "saldo",
    }
    faltan = {"fecha", "ruc", "razon_social", "codigo", "cuenta", "saldo"} - {
        renombre.get(h) for h in header
    }
    if faltan:
        raise ValueError(
            f"{zip_path.name}: encabezado EEFF inesperado {header} (faltan {faltan})"
        )

    partes, cuentas, entidades = [], {}, {}
    for df in pd.read_csv(
        texto,
        sep=sep,
        header=None,
        names=[renombre.get(h, h) for h in header],
        dtype=str,
        keep_default_na=False,
        quotechar='"',
        chunksize=chunksize,
    ):
        df["saldo"] = _a_numero(df["saldo"])
        df["codigo"] = df["codigo"].str.strip()
        for codigo, cuenta in (
            df[["codigo", "cuenta"]].drop_duplicates("codigo").itertuples(index=False)
        ):
            cuentas.setdefault(codigo, cuenta.strip())
        df = df[df["saldo"].fillna(0) != 0]
        for ruc, razon, seg in (
            df[["ruc", "razon_social", "segmento"]]
            .drop_duplicates("ruc")
            .itertuples(index=False)
        ):
            ruc = _ruc_str(ruc)
            if ruc not in entidades:
                tipo = _tipo_entidad(
                    ruc,
                    "MUTUALISTA" if "MUTUALISTA" in str(seg).upper() else "COOPERATIVA",
                )
                entidades[ruc] = resolver_entidad_seps(razon, ruc, tipo)
        partes.append(df[["fecha", "ruc", "razon_social", "codigo", "saldo"]])
    for _ in gen:  # consumir el resto del zip (un solo miembro esperado)
        pass

    df = pd.concat(partes, ignore_index=True)
    df["fecha"] = pd.to_datetime(
        df["fecha"].str.strip().str.strip('"'), format="%Y-%m-%d"
    ).dt.date.map(month_end_date)
    df["ruc"] = df["ruc"].map(_ruc_str)
    df["banco_codigo"] = "BCE_" + df["ruc"]
    df["banco"] = df["razon_social"].str.strip() + " (" + df["ruc"] + ")"
    df["source_file"] = zip_path.name
    es_pyg = df["codigo"].str[0].isin(["4", "5"])
    base_cols = ["fecha", "banco", "banco_codigo", "codigo"]
    # misma forma que parse_boletin_file(): source_hash va aparte (load_raw_boletin)
    balance = df.loc[~es_pyg, base_cols + ["saldo", "source_file"]].rename(
        columns={"saldo": "saldo_usd"}
    )
    pyg = df.loc[es_pyg, base_cols + ["saldo", "source_file"]].rename(
        columns={"saldo": "valor_usd"}
    )

    from benchmark_bancos.transform.parse_boletin import (
        _SECCION_POR_DIGITO,
        _codigo_padre,
    )

    cuentas_df = pd.DataFrame(
        [
            {
                "reporte": "PYG" if c[0] in "45" else "BALANCE",
                "codigo": c,
                "cuenta": d,
                "nivel": len(c),
                "codigo_padre": _codigo_padre(c),
                "seccion": _SECCION_POR_DIGITO.get(c[0]),
                "grupo_met": None,
            }
            for c, d in sorted(cuentas.items())
        ]
    )
    return {
        "balance": balance.reset_index(drop=True),
        "pyg": pyg.reset_index(drop=True),
        "cuentas": cuentas_df,
        "entidades": list(entidades.values()),
        "source_hash": source_hash,
    }


class _IterStream(io.RawIOBase):
    """Adapta un iterador de bytes (miembro de stream_unzip) a un file-like legible."""

    def __init__(self, it):
        self._it = iter(it)
        self._buf = b""

    def readable(self):
        return True

    def readinto(self, b):
        while not self._buf:
            try:
                self._buf = next(self._it)
            except StopIteration:
                return 0
        n = min(len(b), len(self._buf))
        b[:n] = self._buf[:n]
        self._buf = self._buf[n:]
        return n
