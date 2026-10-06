"""
Parser del Boletín Financiero Mensual (Superbancos) -- hojas BALANCE y PYG.

Formato ancho: fila 8 es el encabezado real (CÓDIGO, CUENTA, 23 bancos individuales +
9 columnas de agregado que hay que excluir -- BOLETIN_AGGREGATE_COLUMNS). Plan de cuentas
jerárquico del Catálogo Único de Cuentas: código de 1 dígito (rollup de sección), 2
dígitos (grupo), 4 dígitos (cuenta), 6 dígitos (subcuenta). Filas sin CÓDIGO (encabezado
de sección tipo "ACTIVO"/"PASIVOS", o subtotales tipo "MARGEN NETO INTERESES" en PYG) no
son cuentas reales -- se descartan (ver docs/fuentes_datos.md sección 3).

Valores en miles de USD (confirmado en el encabezado del archivo real,
"(en miles de dólares)") -- se multiplica x1000 al cargar.

seccion se deriva del primer dígito del código (1 activo, 2 pasivo, 3 patrimonio,
4 gastos, 5 ingresos) -- no de la posición de la fila, porque en PYG los códigos 4 y 5
están intercalados por margen, no en bloques separados.
"""

import re
import unicodedata
from pathlib import Path

import openpyxl
import pandas as pd

from benchmark_bancos.transform.banco_matching import resolver_banco_codigo
from benchmark_bancos.transform.common import extract_single_xlsx, normalize_text


def _strip_accents(text: str) -> str:
    return "".join(
        c for c in unicodedata.normalize("NFKD", text) if not unicodedata.combining(c)
    )


def _normalize_col(text: str) -> str:
    """Comparación robusta de nombres de columna de agregado entre plantillas de
    distintos años: mayúsculas, sin acentos, espacios internos colapsados a uno solo."""
    return re.sub(r"\s+", " ", _strip_accents(text.strip().upper()))


# Nombres reales vistos en distintos años del boletín -- varían en tilde/espaciado
# (ej. 'BANCA MÚLTIPLE' 2026 vs 'BANCA MULTIPLE' 2021) y algunos cambiaron de nombre
# entre plantillas (ej. 'BANCOS PRIVADOS MICROCRÉDITO' 2026 vs 'BANCOS PRIVADOS DE
# MICROEMPRESA' 2021) -- se guarda ya normalizado, ver _normalize_col.
BOLETIN_AGGREGATE_COLUMNS = {
    _normalize_col(c)
    for c in (
        "BANCOS PRIVADOS GRANDES",
        "BANCOS PRIVADOS MEDIANOS",
        "BANCOS PRIVADOS PEQUEÑOS",
        "TOTAL BANCOS PRIVADOS",
        "BANCOS PRIVADOS COMERCIALES",
        "BANCOS PRIVADOS CONSUMO",
        "BANCOS PRIVADOS VIVIENDA",
        "BANCOS PRIVADOS MICROCRÉDITO",
        "BANCA MÚLTIPLE",
        "BANCOS PRIVADOS DE MICROEMPRESA",
        "BANCA MULTIPLE",
        "BANCOS PRIVADOS DE MICROCRÉDITO",
    )
}

_CODIGO_VALIDO = re.compile(r"^\d+$")

_SECCION_POR_DIGITO = {
    "1": "ACTIVO",
    "2": "PASIVO",
    "3": "PATRIMONIO",
    "4": "GASTOS",
    "5": "INGRESOS",
    "6": "CONTINGENTE",
    "7": "CUENTAS_DE_ORDEN",
}


def _codigo_padre(codigo: str) -> str | None:
    if len(codigo) <= 1:
        return None
    if len(codigo) == 2:
        return codigo[0]
    return codigo[:-2]


def _find_header_row(ws) -> tuple[int, dict[str, int]]:
    """Boletines antiguos (~2021-2022) usan 'CODIGO' sin tilde -- comparación
    insensible a acentos para cubrir ambas variantes de plantilla. Se exige que la fila
    tenga TANTO 'CODIGO' COMO 'CUENTA': algunos archivos antiguos traen una fila previa
    con solo 'CÓDIGO' (sin 'CUENTA') que sería un falso positivo."""
    for row_idx, row in enumerate(
        ws.iter_rows(min_row=1, max_row=10, values_only=True), start=1
    ):
        cols = {}
        for i, v in enumerate(row):
            if isinstance(v, str) and v.strip():
                key = (
                    "CÓDIGO"
                    if _strip_accents(v.strip().upper()) == "CODIGO"
                    else v.strip()
                )
                cols[key] = i
        if "CÓDIGO" in cols and "CUENTA" in cols:
            return row_idx, cols
    raise ValueError(
        "No se encontró la fila de encabezado (columnas 'CÓDIGO'+'CUENTA')"
    )


def _parse_hoja(ws, cuenta_col_key: str = "CUENTA") -> pd.DataFrame:
    header_row_idx, cols = _find_header_row(ws)
    codigo_idx, cuenta_idx = cols["CÓDIGO"], cols[cuenta_col_key]
    banco_cols = {
        banco: idx
        for banco, idx in cols.items()
        if idx not in (codigo_idx, cuenta_idx)
        and _normalize_col(banco) not in BOLETIN_AGGREGATE_COLUMNS
    }

    records = []
    for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
        codigo = row[codigo_idx] if codigo_idx < len(row) else None
        codigo = str(codigo).strip() if codigo is not None else None
        if not codigo or not _CODIGO_VALIDO.match(codigo):
            continue  # encabezado de sección o subtotal (sin código real), no una cuenta
        cuenta = normalize_text(row[cuenta_idx] if cuenta_idx < len(row) else None)
        for banco, idx in banco_cols.items():
            valor = row[idx] if idx < len(row) else None
            if valor is None:
                continue
            records.append(
                {
                    "codigo": codigo,
                    "cuenta": cuenta,
                    "banco": banco,
                    "valor_miles": float(valor),
                }
            )
    return pd.DataFrame.from_records(records)


def _parse_met(ws) -> dict[str, str]:
    """MET reagrupa códigos de BALANCE en categorías funcionales -- NO es una partición
    limpia (un mismo código puede aparecer en más de un grupo, ej. '11 FONDOS
    DISPONIBLES' cae en 'ACTIVOS LIQUIDOS' y también en 'ACTIVOS IMPRODUCTIVOS BRUTOS').
    Se guarda el PRIMER grupo encontrado por código -- conveniencia documentada en
    docs/metricas_financieras.md, no una verdad exacta de partición."""
    header_row_idx, cols = _find_header_row(ws)
    codigo_idx = cols["CÓDIGO"]
    grupo_actual = None
    grupo_por_codigo: dict[str, str] = {}
    for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
        codigo = row[codigo_idx] if codigo_idx < len(row) else None
        cuenta = row[codigo_idx + 1] if codigo_idx + 1 < len(row) else None
        codigo = str(codigo).strip() if codigo is not None else None
        cuenta = normalize_text(cuenta)
        if not codigo or not _CODIGO_VALIDO.match(codigo):
            if cuenta and not (cuenta.startswith("TOTAL ")):
                grupo_actual = cuenta  # fila de encabezado de grupo funcional
            continue
        if codigo not in grupo_por_codigo:
            grupo_por_codigo[codigo] = grupo_actual
    return grupo_por_codigo


def parse_boletin_file(
    source_path: Path, extract_dir: Path, fecha
) -> dict[str, pd.DataFrame]:
    """Devuelve {'balance': df, 'pyg': df, 'cuentas': df} -- 'cuentas' es el plan de
    cuentas (con seccion/nivel/codigo_padre/grupo_met) descubierto en este archivo, para
    poblar/actualizar dim_cuenta_contable; 'balance'/'pyg' son los valores por banco ya
    con banco_codigo resuelto y x1000 aplicado."""
    xlsx_path = (
        extract_single_xlsx(source_path, extract_dir)
        if source_path.suffix.lower() == ".zip"
        else source_path
    )
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)

    balance = _parse_hoja(wb["BALANCE"])
    pyg = _parse_hoja(wb["PYG"])
    grupo_met = _parse_met(wb["MET"]) if "MET" in wb.sheetnames else {}
    wb.close()

    cuentas = []
    for reporte, df in (("BALANCE", balance), ("PYG", pyg)):
        for codigo, cuenta in (
            df[["codigo", "cuenta"]].drop_duplicates().itertuples(index=False)
        ):
            cuentas.append(
                {
                    "reporte": reporte,
                    "codigo": codigo,
                    "cuenta": cuenta,
                    "nivel": len(codigo),
                    "codigo_padre": _codigo_padre(codigo),
                    "seccion": _SECCION_POR_DIGITO.get(codigo[0]),
                    "grupo_met": (
                        grupo_met.get(codigo) if reporte == "BALANCE" else None
                    ),
                }
            )
    cuentas_df = pd.DataFrame.from_records(cuentas)

    for df, valor_col, out_col in (
        (balance, "valor_miles", "saldo_usd"),
        (pyg, "valor_miles", "valor_usd"),
    ):
        df["banco_codigo"] = df["banco"].apply(
            lambda b: resolver_banco_codigo(b, "BOLETIN")
        )
        df["fecha"] = fecha
        df[out_col] = df[valor_col] * 1000.0
        df["source_file"] = source_path.name
        df.drop(columns=[valor_col], inplace=True)

    return {"balance": balance, "pyg": pyg, "cuentas": cuentas_df}
