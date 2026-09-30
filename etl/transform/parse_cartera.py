"""
Parser del reporte de cartera (colocaciones) de Superbancos.

El nombre del archivo agrupa uno o más tipo_credito (ej. "Cartera de Vivienda" trae
DOS hojas BASE: inmobiliario y vivienda de interés público), así que tipo_credito se
deriva del nombre de cada hoja 'BASE ...', no del nombre del archivo. Cada hoja trae
la serie mensual completa del año en formato ancho (una columna por estado de cartera:
por vencer / no devenga / vencida), que se pasa a formato largo para staging.cartera.
"""

import unicodedata
from pathlib import Path

import openpyxl
import pandas as pd

from etl.config import TIPO_CREDITO_KEYWORDS, TIPOS_ENTIDAD_CAPCOL
from etl.transform.banco_matching import resolver_banco_codigo
from etl.transform.common import (
    extract_single_xlsx,
    find_base_sheets,
    month_end_date,
    normalize_banco,
    normalize_provincia,
    normalize_text,
    region_for_provincia,
    sha256_file,
)

ESTADO_COLUMNS = {
    "POR VENCER": "por_vencer",
    "NO DEVENGA INTERESES": "no_devenga_intereses",
    "VENCIDA": "vencida",
}


def _strip_accents(text: str) -> str:
    return "".join(
        c for c in unicodedata.normalize("NFKD", text) if not unicodedata.combining(c)
    )


def tipo_credito_from_sheet_name(sheet_name: str) -> str:
    normalized = _strip_accents(sheet_name.lower())
    for keyword, canonical in TIPO_CREDITO_KEYWORDS.items():
        if keyword in normalized:
            return canonical
    raise ValueError(f"No se pudo determinar tipo_credito para la hoja '{sheet_name}'")


def _find_header_row(ws):
    for row_idx, row in enumerate(
        ws.iter_rows(min_row=1, max_row=10, values_only=True), start=1
    ):
        if any(isinstance(v, str) and v.strip().upper() == "FECHA" for v in row):
            return row_idx, {
                v.strip().upper(): i
                for i, v in enumerate(row)
                if isinstance(v, str) and v.strip()
            }
    raise ValueError(
        "No se encontró la fila de encabezado (columna 'FECHA') en la hoja BASE"
    )


def _parse_sheet(ws, tipo_credito: str, tipo_entidad: str) -> list[dict]:
    header_row_idx, cols = _find_header_row(ws)

    def get(row, key):
        idx = cols.get(key)
        return row[idx] if idx is not None and idx < len(row) else None

    records = []
    for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
        fecha = get(row, "FECHA")
        if fecha is None:
            continue
        provincia = normalize_provincia(get(row, "PROVINCIA"))
        banco = normalize_banco(get(row, "ENTIDAD"))
        base = {
            "fecha": month_end_date(fecha),
            "tipo_entidad": tipo_entidad,
            "banco": banco,
            "banco_codigo": resolver_banco_codigo(banco, "CAPCOL"),
            "region": region_for_provincia(provincia),
            "provincia": provincia,
            "canton": normalize_text(get(row, "CANTON")),
            "tipo_credito": tipo_credito,
        }
        for col_name, estado in ESTADO_COLUMNS.items():
            saldo = get(row, col_name)
            if saldo is None:
                continue
            records.append({**base, "estado_cartera": estado, "saldo": float(saldo)})
    return records


def parse_cartera_file(
    source_path: Path, extract_dir: Path, tipo_entidad: str = "BANCO PRIVADO"
) -> pd.DataFrame:
    """tipo_entidad lo decide el caller según el sub-portal CAPCOL de origen
    (etl.config.CAPCOL_PORTALES) -- el archivo no lo trae y no se infiere del dato."""
    if tipo_entidad not in TIPOS_ENTIDAD_CAPCOL:
        raise ValueError(
            f"tipo_entidad inválido para CAPCOL: {tipo_entidad!r} "
            f"(debe ser uno de {sorted(TIPOS_ENTIDAD_CAPCOL)})"
        )
    source_hash = sha256_file(source_path)
    xlsx_path = (
        extract_single_xlsx(source_path, extract_dir)
        if source_path.suffix.lower() == ".zip"
        else source_path
    )

    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    records = []
    for sheet_name in find_base_sheets(wb):
        tipo_credito = tipo_credito_from_sheet_name(sheet_name)
        records.extend(_parse_sheet(wb[sheet_name], tipo_credito, tipo_entidad))
    wb.close()

    df = pd.DataFrame.from_records(records)
    df["source_file"] = source_path.name
    df["source_hash"] = source_hash

    # El origen trae detalle a nivel de oficina/cuenta contable: varias filas pueden
    # compartir (fecha, banco, canton, tipo_credito, estado_cartera) y deben sumarse,
    # no sobrescribirse, para no perder saldo al cargar a staging.
    key_cols = [
        "fecha",
        "tipo_entidad",
        "banco",
        "banco_codigo",
        "region",
        "provincia",
        "canton",
        "tipo_credito",
        "estado_cartera",
    ]
    df = df.groupby(key_cols, dropna=False, as_index=False).agg(
        saldo=("saldo", "sum"),
        source_file=("source_file", "first"),
        source_hash=("source_hash", "first"),
    )
    return df
