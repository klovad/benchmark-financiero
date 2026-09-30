"""
Parser del reporte de depósitos (captaciones) de Superbancos.

A diferencia de cartera, un solo archivo por año cubre todos los tipos de depósito;
'TIPO DE DEPOSITO' ya viene como columna en la hoja 'BASE ...' (más granular que las
categorías de la ficha metodológica: incluye buckets de plazo por rango de días).
"""

from pathlib import Path

import openpyxl
import pandas as pd

from etl.config import TIPOS_ENTIDAD_CAPCOL
from etl.transform.banco_matching import resolver_banco_codigo
from etl.transform.categoria_deposito_matching import resolver_categoria_deposito
from etl.transform.common import (
    extract_single_xlsx,
    find_base_sheets,
    month_end_date,
    normalize_banco,
    normalize_provincia,
    normalize_text,
    region_for_provincia,
    sha256_file,
)


def _find_header_row(ws):
    for row_idx, row in enumerate(
        ws.iter_rows(min_row=1, max_row=10, values_only=True), start=1
    ):
        if any(isinstance(v, str) and v.strip().upper() == "FECHA" for v in row):
            return row_idx, {
                v.strip().upper(): i
                for i, v in enumerate(row)
                if isinstance(v, str) and v.strip()
            }
    raise ValueError(
        "No se encontró la fila de encabezado (columna 'FECHA') en la hoja BASE"
    )


def parse_depositos_file(
    source_path: Path, extract_dir: Path, tipo_entidad: str = "BANCO PRIVADO"
) -> pd.DataFrame:
    """tipo_entidad lo decide el caller según el sub-portal CAPCOL de origen
    (etl.config.CAPCOL_PORTALES) -- el archivo no lo trae y no se infiere del dato."""
    if tipo_entidad not in TIPOS_ENTIDAD_CAPCOL:
        raise ValueError(
            f"tipo_entidad inválido para CAPCOL: {tipo_entidad!r} "
            f"(debe ser uno de {sorted(TIPOS_ENTIDAD_CAPCOL)})"
        )
    source_hash = sha256_file(source_path)
    xlsx_path = (
        extract_single_xlsx(source_path, extract_dir)
        if source_path.suffix.lower() == ".zip"
        else source_path
    )

    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    sheet_name = find_base_sheets(wb)[0]
    ws = wb[sheet_name]
    header_row_idx, cols = _find_header_row(ws)

    def get(row, key):
        idx = cols.get(key)
        return row[idx] if idx is not None and idx < len(row) else None

    records = []
    for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
        fecha = get(row, "FECHA")
        if fecha is None:
            continue
        saldo = get(row, "SALDO")
        if saldo is None:
            continue
        provincia = normalize_provincia(get(row, "PROVINCIA"))
        numero_cuentas = get(row, "NUMERO DE CUENTAS")
        numero_clientes = get(row, "NUMERO DE CLIENTES")
        banco = normalize_banco(get(row, "ENTIDAD"))
        tipo_deposito = normalize_text(get(row, "TIPO DE DEPOSITO"))
        categoria_deposito, plazo_dias_desde, plazo_dias_hasta = (
            resolver_categoria_deposito(tipo_deposito)
        )
        records.append(
            {
                "fecha": month_end_date(fecha),
                "tipo_entidad": tipo_entidad,
                "banco": banco,
                "banco_codigo": resolver_banco_codigo(banco, "CAPCOL"),
                # Derivado siempre de PROVINCIA_REGION (no de la columna REGION del archivo
                # fuente): se confirmó que esa columna trae valores inconsistentes con
                # cartera para al menos una provincia (MORONA SANTIAGO: "AMAZONICA" en
                # depositos vs. "ORIENTE", el mapeo canónico, en cartera) -- marts.dim_canton
                # llegó a tener 2 regiones distintas para la misma provincia por esto.
                "region": region_for_provincia(provincia),
                "provincia": provincia,
                "canton": normalize_text(get(row, "CANTON")),
                "tipo_deposito": tipo_deposito,
                "categoria_deposito": categoria_deposito,
                "plazo_dias_desde": plazo_dias_desde,
                "plazo_dias_hasta": plazo_dias_hasta,
                "saldo": float(saldo),
                "numero_cuentas": (
                    int(numero_cuentas) if numero_cuentas is not None else None
                ),
                "numero_clientes": (
                    int(numero_clientes) if numero_clientes is not None else None
                ),
            }
        )

    wb.close()
    df = pd.DataFrame.from_records(records)
    df["source_file"] = source_path.name
    df["source_hash"] = source_hash

    # El origen trae una fila por cuenta contable (columna CUENTA, no conservada aquí);
    # varias cuentas comparten (fecha, banco, canton, tipo_deposito) y deben sumarse,
    # no sobrescribirse, para no perder saldo al cargar a staging.
    key_cols = [
        "fecha",
        "tipo_entidad",
        "banco",
        "banco_codigo",
        "region",
        "provincia",
        "canton",
        "tipo_deposito",
    ]
    df = df.groupby(key_cols, dropna=False, as_index=False).agg(
        saldo=("saldo", "sum"),
        numero_cuentas=("numero_cuentas", "sum"),
        numero_clientes=("numero_clientes", "sum"),
        categoria_deposito=("categoria_deposito", "first"),
        plazo_dias_desde=("plazo_dias_desde", "first"),
        plazo_dias_hasta=("plazo_dias_hasta", "first"),
        source_file=("source_file", "first"),
        source_hash=("source_hash", "first"),
    )
    return df
